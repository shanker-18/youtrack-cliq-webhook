# pylint: disable=too-many-return-statements,too-many-nested-blocks,global-statement,too-many-lines
"""
NA IBP Planning Application — Validation Calculations Engine
============================================================
Authoritative calculation engine for both Consumption and Shipment validation:
- Uses services.kv_calendar.KenvueCalendar for pure 4-4-5 fiscal math (zero Excel dump dependency)
- Price Index lookups using openpyxl & fallback maps
- PostgreSQL Live Building Block additions (11 official POS adjustment blocks)
- Statistical 2-sigma outlier detection on Build metrics
- 100% Dynamic year and month rollover (no hardcoded years or cutoffs)
- Pure Python implementation (zero pandas dependency)
"""

import os
import re
import math
import statistics
from datetime import date, datetime
from typing import Optional, Dict, List, Tuple, Any, Set
import psycopg2

from config import config
from database.connection import get_db_cursor
from services.kv_calendar import KenvueCalendar, MONTH_NAMES, MONTH_ABBR
from utils.constants import DRIVER_NET_INNOVATION, DRIVER_NET_PRICE
from utils.logger import get_logger

logger = get_logger(__name__)

MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]

MONTH_NUMBER_TO_3LETTER = {
    1: "JAN", 2: "FEB", 3: "MAR", 4: "APR", 5: "MAY", 6: "JUN",
    7: "JUL", 8: "AUG", 9: "SEP", 10: "OCT", 11: "NOV", 12: "DEC"
}

MONTH_NAME_MAP = {
    "JAN": "JAN", "JANUARY": "JAN",
    "FEB": "FEB", "FEBRUARY": "FEB",
    "MAR": "MAR", "MARCH": "MAR",
    "APR": "APR", "APRIL": "APR",
    "MAY": "MAY",
    "JUN": "JUN", "JUNE": "JUN",
    "JUL": "JUL", "JULY": "JUL",
    "AUG": "AUG", "AUGUST": "AUG",
    "SEP": "SEP", "SEPTEMBER": "SEP",
    "OCT": "OCT", "OCTOBER": "OCT",
    "NOV": "NOV", "NOVEMBER": "NOV",
    "DEC": "DEC", "DECEMBER": "DEC",
}

# 11 Official POS Adjustment Building Blocks
ALLOWED_BUILDING_BLOCKS = [
    "Base Trend",
    "Season",
    "Competition",
    "Distribution",
    DRIVER_NET_INNOVATION,
    "Discontinuations",
    "Renovation",
    "Media",
    "Trade",
    DRIVER_NET_PRICE,
    "Club",
]

RAW_BB_MAP = {
    "BASE TREND": "Base Trend",
    "SEASON": "Season",
    "COMPETITION": "Competition",
    "DISTRIBUTION": "Distribution",
    "INNOVATION": DRIVER_NET_INNOVATION,
    "CANNIBALIZATION": DRIVER_NET_INNOVATION,
    "NET INNOVATION": DRIVER_NET_INNOVATION,
    "DISCONTINUATIONS": "Discontinuations",
    "RENOVATION": "Renovation",
    "MEDIA": "Media",
    "TRADE": "Trade",
    "PRICE": DRIVER_NET_PRICE,
    "ELASTICITY": DRIVER_NET_PRICE,
    "NET PRICE": DRIVER_NET_PRICE,
    "CLUB": "Club",
}

_CACHED_PRICE_INDEX_MAP = None


def get_current_kv_year() -> int:
    """Helper to dynamically determine active Kenvue plan year."""
    env_yr = os.getenv("IBP_PLAN_YEAR")
    if env_yr and env_yr.isdigit():
        return int(env_yr)
    return KenvueCalendar.get_date_info(date.today()).kv_year


def normalize_text(val: Any) -> str:
    """Normalize string consistently across Excel, PostgreSQL, and Snowflake."""
    if val is None:
        return ""
    s = str(val).replace("\xa0", " ").strip().upper()
    if s in ["NAN", "NONE", "NULL", "EMPTY", "N/A", "<NA>"]:
        return ""
    return re.sub(r"\s+", " ", s)


def _normalize_numeric_month(val: Any) -> str:
    try:
        n = int(val)
        if 1 <= n <= 12:
            return MONTH_NUMBER_TO_3LETTER[n]
    except (ValueError, TypeError, OverflowError):
        pass
    return ""


def _parse_month_token(s: str) -> str:
    for t in re.findall(r"[A-Z0-9]+", s):
        if t in MONTH_NAME_MAP:
            return MONTH_NAME_MAP[t]
    m_iso = re.search(r"(\d{4})[-/](\d{1,2})", s)
    if m_iso:
        mo = int(m_iso.group(2))
        if 1 <= mo <= 12:
            return MONTH_NUMBER_TO_3LETTER[mo]
    return ""


def _normalize_string_month(s: str) -> str:
    s_clean = s.strip().upper()
    if not s_clean or s_clean in ("NAN", "NONE", "NULL", "N/A"):
        return ""
    if s_clean.isdigit():
        return _normalize_numeric_month(s_clean)
    return _parse_month_token(s_clean)


def normalize_month_3letter(val: Any) -> str:
    """Normalizes any month representation into standard 3-letter uppercase ('JAN'..'DEC')."""
    if val is None:
        return ""
    if isinstance(val, (date, datetime)):
        return MONTH_NUMBER_TO_3LETTER.get(val.month, "")
    if isinstance(val, (int, float)):
        return _normalize_numeric_month(val)
    return _normalize_string_month(str(val))


# ─── 1. Kenvue Calendar 4-4-5 Fiscal Ranges via services.kv_calendar ─────────

def get_kv_month_ranges(start_year: Optional[int] = None, end_year: Optional[int] = None) -> List[dict]:
    """Generates Kenvue Fiscal Month date ranges dynamically using KenvueCalendar (Rolling Window)."""
    curr_yr = get_current_kv_year()
    s_yr = start_year if start_year is not None else (curr_yr - 5)
    e_yr = end_year if end_year is not None else (curr_yr + 1)

    ranges = []
    for yr in range(s_yr, e_yr + 1):
        for m in range(1, 13):
            c_start, c_end = KenvueCalendar.get_cycle_bounds(yr, m)
            ranges.append({
                "year": yr,
                "y_str": str(yr),
                "m_nbr": m,
                "kv_month": f"{yr}-{m:02d}",
                "start": c_start.isoformat(),
                "end": c_end.isoformat(),
            })
    return ranges


def map_date_to_kv_calendar(d_val: Any) -> Optional[dict]:
    """Maps date value to Kenvue Fiscal Calendar properties using KenvueCalendar."""
    if not d_val:
        return None

    dt_obj = None
    if isinstance(d_val, (date, datetime)):
        dt_obj = d_val if isinstance(d_val, date) else d_val.date()
    else:
        d_clean = str(d_val).strip()[:10]
        try:
            dt_obj = date.fromisoformat(d_clean)
        except (ValueError, TypeError):
            for fmt in ["%d-%b-%y", "%d-%b-%Y", "%m/%d/%Y", "%Y/%m/%d"]:
                try:
                    dt_obj = datetime.strptime(d_clean, fmt).date()
                    break
                except (ValueError, TypeError):
                    pass

    if not dt_obj:
        return None

    info = KenvueCalendar.get_date_info(dt_obj)
    m_idx = info.kv_month_num - 1
    m_name = MONTHS[m_idx]
    y_str = str(info.kv_year)

    return {
        "m_idx": m_idx,
        "m_nbr": info.kv_month_num,
        "m_name": m_name,
        "y_str": y_str,
        "quarter": f"Q{info.kv_quarter}",
        "kv_wk_id": f"{y_str}{info.kv_week_num:02d}",
        "kv_mo_id": f"{y_str}{info.kv_month_num:02d}",
        "kv_week_beginning": dt_obj.isoformat(),
    }


def _update_latest_year(current_yr: Optional[str], candidate_yr: str) -> str:
    if not current_yr or candidate_yr > current_yr:
        return candidate_yr
    return current_yr


def _record_row_calendar_week(
    r: Dict[str, Any], avail_weeks_by_month: dict, latest_yr: Optional[str]
) -> Optional[str]:
    d_str = str(r.get("GLOBAL_DATE_SHORT_DESC") or "").strip()
    kv_info = map_date_to_kv_calendar(d_str) if d_str else None
    if kv_info:
        y_s = kv_info["y_str"]
        m_n = kv_info["m_nbr"]
        avail_weeks_by_month.setdefault((y_s, m_n), set()).add(d_str or kv_info["kv_wk_id"])
        return _update_latest_year(latest_yr, y_s)

    kv_m_str = str(r.get("KV_MONTH") or "").strip()
    if "-" in kv_m_str:
        parts = kv_m_str.split("-")
        y_s = parts[0]
        m_n = int(parts[1])
        avail_weeks_by_month.setdefault((y_s, m_n), set()).add(d_str or kv_m_str)
        return _update_latest_year(latest_yr, y_s)

    return latest_yr


def _record_row_completed_month(
    r: Dict[str, Any], avail_weeks_by_month: dict, latest_yr: Optional[str]
) -> Optional[str]:
    if r.get("is_complete") is not True:
        return latest_yr
    kv_m = str(r.get("KV_MONTH") or "")
    y_s = str(r.get("YEAR") or (kv_m.split("-", maxsplit=1)[0] if "-" in kv_m else ""))
    m_n = int(r.get("MONTHNUMBER") or (kv_m.split("-", maxsplit=1)[1] if "-" in kv_m else 0))
    if y_s and m_n:
        avail_weeks_by_month.setdefault((y_s, m_n), set()).update({f"wk_{i}" for i in range(1, 6)})
        return _update_latest_year(latest_yr, y_s)
    return latest_yr


def _collect_weeks_by_month(rows: List[Dict[str, Any]]) -> Tuple[dict, Optional[str]]:
    avail_weeks_by_month: dict = {}
    latest_yr_found: Optional[str] = None

    for r in rows:
        latest_yr_found = _record_row_calendar_week(r, avail_weeks_by_month, latest_yr_found)
        latest_yr_found = _record_row_completed_month(r, avail_weeks_by_month, latest_yr_found)

    return avail_weeks_by_month, latest_yr_found


def _evaluate_month_completeness_details(
    avail_weeks_by_month: dict, target_yr: str, target_yr_int: int
) -> Tuple[dict, int]:
    month_details = {}
    latest_complete_m_nbr = 0

    for m_n in range(1, 13):
        req_wks = KenvueCalendar.get_month_week_count(target_yr_int, m_n)
        avail_set = avail_weeks_by_month.get((target_yr, m_n), set())
        avail_wks = len(avail_set)
        is_comp = (avail_wks >= req_wks) if req_wks > 0 else False

        month_details[m_n] = {
            "required_wks": req_wks,
            "available_wks": avail_wks,
            "is_complete": is_comp,
        }
        if is_comp and m_n == latest_complete_m_nbr + 1:
            latest_complete_m_nbr = m_n

    if latest_complete_m_nbr == 0:
        for m_n in range(1, 13):
            if month_details[m_n]["is_complete"]:
                latest_complete_m_nbr = m_n

    return month_details, latest_complete_m_nbr


def get_kv_month_completeness_status(rows: List[Dict[str, Any]], target_year: Optional[str] = None) -> dict:
    """
    Evaluates month completeness dynamically without hardcoded fallback years.
    """
    avail_weeks_by_month, latest_yr_found = _collect_weeks_by_month(rows)

    curr_kv_yr = get_current_kv_year()
    target_yr = str(target_year) if target_year else (latest_yr_found or str(curr_kv_yr))
    target_yr_int = int(target_yr) if target_yr.isdigit() else curr_kv_yr

    month_details, latest_complete_m_nbr = _evaluate_month_completeness_details(
        avail_weeks_by_month, target_yr, target_yr_int
    )

    if latest_complete_m_nbr == 0:
        latest_complete_m_idx = -1
        m_name = "NONE"
        ytd_s = [0, 0]
        ytg_s = [0, 12]
    else:
        latest_complete_m_idx = latest_complete_m_nbr - 1
        m_name = MONTHS[latest_complete_m_idx] if 0 <= latest_complete_m_idx < 12 else "DEC"
        ytd_s = [0, latest_complete_m_nbr]
        ytg_s = [latest_complete_m_nbr, 12]

    return {
        "latest_year": target_yr,
        "latest_complete_m_nbr": latest_complete_m_nbr,
        "latest_complete_m_idx": latest_complete_m_idx,
        "latest_complete_m_name": m_name,
        "ytd_slice": ytd_s,
        "ytg_slice": ytg_s,
        "month_details": month_details,
    }


# ─── 2. Price Index Reference & Factory POS Calculations ─────────────────────

def _find_price_index_headers(rows: list) -> Tuple[Optional[int], Optional[int], Optional[int], int]:
    hdr_idx = 0
    for idx, r in enumerate(rows[:5]):
        r_strs = [str(c).upper() for c in r if c is not None]
        if any("MNTH" in c or "MONTH" in c for c in r_strs) and any("MODEL" in c for c in r_strs):
            hdr_idx = idx
            break

    hdr = [str(c).strip().upper() if c else "" for c in rows[hdr_idx]]
    mnth_i = next((i for i, c in enumerate(hdr) if "MNTH" in c or "MONTH" in c), None)
    model_i = next((i for i, c in enumerate(hdr) if "MODEL" in c), None)
    index_i = next((i for i, c in enumerate(hdr) if "INDEX" in c), None)
    return mnth_i, model_i, index_i, hdr_idx


def _parse_price_index_sheet_rows(rows: list, lookup: dict):
    if len(rows) < 2:
        return
    mnth_i, model_i, index_i, hdr_idx = _find_price_index_headers(rows)
    if mnth_i is None or model_i is None or index_i is None:
        return

    for data_row in rows[hdr_idx + 1:]:
        if len(data_row) > max(mnth_i, model_i, index_i):
            norm_m = normalize_month_3letter(data_row[mnth_i])
            norm_mod = normalize_text(data_row[model_i])
            idx_raw = data_row[index_i]
            if norm_m and norm_mod and idx_raw is not None:
                try:
                    lookup[(norm_m, norm_mod)] = float(idx_raw)
                except (ValueError, TypeError):
                    pass


def _load_price_index_from_excel(candidate_paths: list) -> dict:
    lookup = {}
    try:
        import openpyxl
        for cand in candidate_paths:
            if cand and os.path.exists(cand):
                try:
                    wb = openpyxl.load_workbook(cand, data_only=True, read_only=True)
                    ws = wb["Price Index"] if "Price Index" in wb.sheetnames else wb.active
                    rows = list(ws.iter_rows(values_only=True))
                    wb.close()
                    _parse_price_index_sheet_rows(rows, lookup)
                    if lookup:
                        break
                except (OSError, IOError, ValueError, KeyError) as e:
                    logger.warning("Notice reading Price Index via openpyxl: %s", e)
    except ImportError:
        pass
    return lookup


def load_price_index_map() -> dict:
    """Loads Price Index multipliers from file or openpyxl into (norm_month, norm_model) -> multiplier."""
    global _CACHED_PRICE_INDEX_MAP
    if _CACHED_PRICE_INDEX_MAP is not None:
        return _CACHED_PRICE_INDEX_MAP

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    candidate_paths = [
        config.PRICE_INDEX_EXCEL_PATH,
        os.path.join(base_dir, "validation", "Price Index.xlsx"),
        os.path.join(base_dir, "data", "Price Index.xlsx"),
    ]

    lookup = _load_price_index_from_excel(candidate_paths)

    # Built-in fallbacks for standard OTC models
    adult_sudafed_indices = {
        "JAN": 0.7600, "FEB": 0.7800, "MAR": 0.7500, "APR": 0.7400,
        "MAY": 0.7300, "JUN": 0.7200, "JUL": 0.7100, "AUG": 0.7000,
        "SEP": 0.7500, "OCT": 0.7700, "NOV": 0.7900, "DEC": 0.8000,
    }
    for m, val in adult_sudafed_indices.items():
        lookup.setdefault((m, "ADULT SUDAFED"), val)

    _CACHED_PRICE_INDEX_MAP = lookup
    return lookup


def get_factory_pos_val(month_num: int, model_name: str | list[str], pos_val: float | None) -> Tuple[float | None, float]:
    """Calculates Factory POS $ = POS $ * Price Index multiplier."""
    if pos_val is None:
        return None, 1.0
    norm_m = normalize_month_3letter(month_num)
    idx_map = load_price_index_map()
    if isinstance(model_name, (list, tuple, set)):
        clean_models = [normalize_text(m) for m in model_name if m]
        indices = [idx_map.get((norm_m, m), 1.0) for m in clean_models if (norm_m, m) in idx_map]
        index_val = (sum(indices) / len(indices)) if indices else 1.0
    else:
        norm_model = normalize_text(model_name)
        index_val = idx_map.get((norm_m, norm_model), 1.0)
    return float(pos_val) * index_val, index_val


# ─── 3. Live PostgreSQL Building Block Planning Values ───────────────────────

def load_live_building_blocks(domain: str = "CONSUMPTION") -> List[Dict[str, Any]]:
    """
    Queries live PostgreSQL planning rows and values for active building blocks.
    Returns list of dicts: [{'model': ..., 'period_month': ..., 'building_block': ..., 'value_in_thousands': ...}].
    """
    table_prefix = "shipment" if domain.upper() == "SHIPMENT" else "consumption"
    sql = f"""
    SELECT
        hm.model_name AS model,
        DATE_TRUNC('month', pv.period_month)::date AS period_month,
        bb.name AS building_block,
        SUM(pv.value_in_thousands) AS value_in_thousands
    FROM public.{table_prefix}_planning_rows pr
    INNER JOIN public.{table_prefix}_building_blocks bb
        ON pr.building_block_id = bb.block_id
    INNER JOIN public.{table_prefix}_planning_values pv
        ON pr.row_id = pv.row_id
    INNER JOIN public.master_models hm
        ON pr.model_id = hm.model_id
    WHERE
        pr.is_deleted = FALSE
        AND bb.is_active = TRUE
        AND hm.is_active = TRUE
    GROUP BY
        hm.model_name,
        DATE_TRUNC('month', pv.period_month),
        bb.name
    ORDER BY
        hm.model_name,
        period_month,
        bb.name;
    """
    try:
        with get_db_cursor(commit=False) as cur:
            cur.execute(sql)
            raw_rows = cur.fetchall()
    except (psycopg2.Error, OSError, RuntimeError) as err:
        logger.warning("Could not query live building blocks: %s", err)
        return []

    result = []
    for r in raw_rows:
        result.append({
            "model": str(r["model"]).strip(),
            "period_month": r["period_month"],
            "building_block": str(r["building_block"]).strip(),
            "value_in_thousands": float(r["value_in_thousands"] or 0.0),
        })
    return result


def _matches_period(p_val: Any, year: int, month_nbr: int) -> bool:
    if not p_val:
        return False
    dt = p_val if isinstance(p_val, date) else date.fromisoformat(str(p_val)[:10])
    return dt.year == year and dt.month == month_nbr


def _matches_consumption_bb_row(r: Dict[str, Any], norm_targets: Set[str], year: int, month_nbr: int) -> bool:
    if normalize_text(r.get("model", "")) not in norm_targets:
        return False
    if normalize_text(r.get("building_block", "")) not in RAW_BB_MAP:
        return False
    return _matches_period(r.get("period_month"), year, month_nbr)


def get_building_block_total_for_period(
    model_name: str | list[str], year: int, month_nbr: int, bb_rows: List[Dict[str, Any]]
) -> float:
    """Returns total building block planning value in thousands ($K) for model(s) in period year-month."""
    if not bb_rows or not model_name:
        return 0.0

    if isinstance(model_name, (list, tuple, set)):
        norm_targets = {normalize_text(m) for m in model_name if m}
    else:
        norm_targets = {normalize_text(model_name)}

    return sum(
        float(r.get("value_in_thousands") or 0.0)
        for r in bb_rows
        if _matches_consumption_bb_row(r, norm_targets, year, month_nbr)
    )


SHIPMENT_ALLOWED_BUILDING_BLOCKS = ["Innovation", "Trade", "Club", "Retailer Inventory"]


def map_shipment_bb_name(raw_name: str) -> Optional[str]:
    """Maps Shipment Building Block PostgreSQL names to standard logical names."""
    if not raw_name:
        return None
    norm = normalize_text(raw_name)
    if "INNOVATION" in norm:
        res = "Innovation"
    elif "TRADE" in norm:
        res = "Trade"
    elif "CLUB" in norm:
        res = "Club"
    elif "INVENTORY" in norm or "RETAILER" in norm:
        res = "Retailer Inventory"
    else:
        res = None
    return res


def _matches_shipment_bb_row(r: Dict[str, Any], norm_targets: Set[str], year: int, month_nbr: int) -> bool:
    if normalize_text(r.get("model", "")) not in norm_targets:
        return False
    if not map_shipment_bb_name(r.get("building_block", "")):
        return False
    return _matches_period(r.get("period_month"), year, month_nbr)


def get_shipment_building_block_total_for_period(
    model_name: str | list[str], year: int, month_nbr: int, bb_rows: List[Dict[str, Any]]
) -> float:
    """Returns total shipment building block planning value in thousands ($K) for model(s) in period year-month."""
    if not bb_rows or not model_name:
        return 0.0

    if isinstance(model_name, (list, tuple, set)):
        norm_targets = {normalize_text(m) for m in model_name if m}
    else:
        norm_targets = {normalize_text(model_name)}

    return sum(
        float(r.get("value_in_thousands") or 0.0)
        for r in bb_rows
        if _matches_shipment_bb_row(r, norm_targets, year, month_nbr)
    )


def _extract_row_year_and_month_idx(r: Dict[str, Any]) -> Tuple[Optional[str], Optional[int]]:
    kv_m_str = str(r.get("KV_MONTH") or "").strip()
    if kv_m_str and "-" in kv_m_str:
        parts = kv_m_str.split("-")
        return parts[0], int(parts[1]) - 1

    d_str = str(r.get("GLOBAL_DATE_SHORT_DESC") or "").strip()
    kv_info = map_date_to_kv_calendar(d_str)
    if kv_info:
        return kv_info["y_str"], kv_info["m_idx"]

    return None, None


def _accumulate_pos_maps(
    r: Dict[str, Any],
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
) -> None:
    y_str, m_idx = _extract_row_year_and_month_idx(r)
    curr_yr = get_current_kv_year()
    # Dynamic rolling 5-year floor
    if not y_str or m_idx is None or (y_str.isdigit() and int(y_str) < (curr_yr - 5)):
        return

    pv = r.get("POS_VALUE") if r.get("POS_VALUE") is not None else r.get("POS_DOLLARS")
    pu = r.get("POS_UNITS")
    key = (str(y_str), m_idx)
    if pv is not None:
        pos_val_map[key] = pos_val_map.get(key, 0.0) + float(pv)
    if pu is not None:
        pos_u_map[key] = pos_u_map.get(key, 0.0) + float(pu)


def _build_factory_pos_map(
    pos_val_map: Dict[Tuple[str, int], float], model_name: str | list[str]
) -> Dict[Tuple[str, int], float]:
    factory_pos_map: Dict[Tuple[str, int], float] = {}
    for (y_str, m_idx), pv_val in pos_val_map.items():
        f_pos, _ = get_factory_pos_val(m_idx + 1, model_name, pv_val)
        factory_pos_map[(str(y_str), m_idx)] = f_pos if f_pos is not None else pv_val
    return factory_pos_map


def extract_consumption_maps(
    rows_cons: List[Dict[str, Any]], model_name: str | list[str]
) -> Tuple[Dict[Tuple[str, int], float], Dict[Tuple[str, int], float], Dict[Tuple[str, int], float]]:
    """
    Extracts factory_pos_map, pos_val_map, pos_u_map from Snowflake consumption rows.
    Returns: (factory_pos_map, pos_val_map, pos_u_map) keyed by (year_str, month_idx 0..11).
    """
    if not rows_cons:
        return {}, {}, {}

    pos_val_map: Dict[Tuple[str, int], float] = {}
    pos_u_map: Dict[Tuple[str, int], float] = {}

    for r in rows_cons:
        _accumulate_pos_maps(r, pos_val_map, pos_u_map)

    factory_pos_map = _build_factory_pos_map(pos_val_map, model_name)
    return factory_pos_map, pos_val_map, pos_u_map


# ─── 4. Full Metrics Matrix Calculators ───────────────────────────────────────

def calc_asp(pos_val: Optional[float], pos_units: Optional[float]) -> Optional[float]:
    if pos_val is None or pos_units is None or pos_units == 0:
        return None
    return pos_val / pos_units


def calc_pct_var(actual: Optional[float], comparison: Optional[float]) -> Optional[float]:
    if actual is None or comparison is None or comparison == 0:
        return None
    return ((actual - comparison) / abs(comparison)) * 100.0


def calc_share(month_val: Optional[float], total_val: Optional[float]) -> Optional[float]:
    if month_val is None or total_val is None or total_val == 0:
        return None
    return (month_val / total_val) * 100.0


def _accumulate_year_consumption_metrics(
    r: Dict[str, Any], data_by_year: Dict[str, Dict[str, List[Optional[float]]]]
) -> None:
    y_str, m_idx = _extract_row_year_and_month_idx(r)
    curr_yr = get_current_kv_year()
    # Dynamic rolling 5-year floor
    if not y_str or m_idx is None or (y_str.isdigit() and int(y_str) < (curr_yr - 5)):
        return

    if y_str not in data_by_year:
        data_by_year[y_str] = {
            "pos_val": [None] * 12,
            "factory_pos": [None] * 12,
            "pos_u": [None] * 12,
        }
    d_dict = data_by_year[y_str]
    pv = r.get("POS_VALUE") or r.get("POS_DOLLARS")
    pu = r.get("POS_UNITS")

    if pv is not None:
        d_dict["pos_val"][m_idx] = (d_dict["pos_val"][m_idx] or 0.0) + float(pv)
    if pu is not None:
        d_dict["pos_u"][m_idx] = (d_dict["pos_u"][m_idx] or 0.0) + float(pu)


def _aggregate_consumption_sf_rows(rows_sf: List[Dict[str, Any]]) -> Dict[str, Dict[str, List[Optional[float]]]]:
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]] = {}
    for r in rows_sf:
        _accumulate_year_consumption_metrics(r, data_by_year)
    return data_by_year


def _extract_ytd_slice(comp_status: Dict[str, Any]) -> slice:
    ytd_s_raw = comp_status.get("ytd_slice")
    if isinstance(ytd_s_raw, (list, tuple)) and len(ytd_s_raw) >= 2:
        return slice(ytd_s_raw[0], ytd_s_raw[1])
    if isinstance(ytd_s_raw, slice):
        return ytd_s_raw
    return slice(0, 0)


def _compute_projected_asp(
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]],
    latest_year: str,
    prev_year: str,
    m_i: int,
    prev_yr_pos: float,
    ytd_slice: slice,
) -> Optional[float]:
    prev_yr_u = data_by_year.get(prev_year, {}).get("pos_u", [None] * 12)[m_i]
    if not prev_yr_u or prev_yr_u <= 0:
        return None

    prev_same_asp = prev_yr_pos / prev_yr_u
    curr_ytd_v = sum(v for v in data_by_year[latest_year]["pos_val"][ytd_slice] if v is not None)
    curr_ytd_u = sum(u for u in data_by_year[latest_year]["pos_u"][ytd_slice] if u is not None)
    prev_ytd_v = sum(v for v in data_by_year.get(prev_year, {}).get("pos_val", [None] * 12)[ytd_slice] if v is not None)
    prev_ytd_u = sum(u for u in data_by_year.get(prev_year, {}).get("pos_u", [None] * 12)[ytd_slice] if u is not None)

    curr_asp_ytd = (curr_ytd_v / curr_ytd_u) if curr_ytd_u > 0 else 0
    prev_asp_ytd = (prev_ytd_v / prev_ytd_u) if prev_ytd_u > 0 else 0
    yoy_asp_pct = ((curr_asp_ytd - prev_asp_ytd) / prev_asp_ytd) if prev_asp_ytd > 0 else 0.0

    return prev_same_asp * (1.0 + yoy_asp_pct)


def _project_consumption_future_months(
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]],
    latest_year: str,
    prev_year: str,
    latest_year_int: int,
    latest_comp_m_nbr: int,
    model_name: str | list[str],
    comp_status: Dict[str, Any],
):
    if latest_comp_m_nbr >= 12:
        return
    bb_rows = load_live_building_blocks(domain="CONSUMPTION")
    ytd_slice = _extract_ytd_slice(comp_status)

    for m_n in range(latest_comp_m_nbr + 1, 13):
        m_i = m_n - 1
        prev_yr_pos = data_by_year.get(prev_year, {}).get("pos_val", [None] * 12)[m_i]
        if prev_yr_pos is None:
            continue
        tot_bb_k = get_building_block_total_for_period(model_name, latest_year_int, m_n, bb_rows)
        calc_pos_usd = prev_yr_pos + (tot_bb_k * 1000.0)
        data_by_year[latest_year]["pos_val"][m_i] = calc_pos_usd

        f_pos, _ = get_factory_pos_val(m_n, model_name, calc_pos_usd)
        data_by_year[latest_year]["factory_pos"][m_i] = f_pos

        proj_asp = _compute_projected_asp(data_by_year, latest_year, prev_year, m_i, prev_yr_pos, ytd_slice)
        if proj_asp and proj_asp > 0:
            data_by_year[latest_year]["pos_u"][m_i] = calc_pos_usd / proj_asp


def _calculate_single_year_consumption_metrics(
    d: Dict[str, List[Optional[float]]],
    prev_yr_dec_pos: Optional[float],
    prev_yr_dec_asp: Optional[float],
) -> Dict[str, Any]:
    asp = [None] * 12
    for i in range(12):
        c_asp = calc_asp(d["pos_val"][i], d["pos_u"][i])
        if c_asp is not None:
            asp[i] = c_asp
        elif i > 0:
            asp[i] = asp[i - 1]
        else:
            asp[i] = prev_yr_dec_asp

    build = [None] * 12
    for i in range(12):
        cur_p = d["pos_val"][i]
        prev_p = d["pos_val"][i - 1] if i > 0 else prev_yr_dec_pos
        if cur_p is not None and prev_p is not None and prev_p != 0:
            build[i] = round(cur_p / prev_p, 2)

    tot_p = sum(v for v in d["pos_val"] if v is not None and not (isinstance(v, float) and math.isnan(v)))
    share = [calc_share(d["pos_val"][i], tot_p) for i in range(12)]

    return {
        "pos_val": d["pos_val"],
        "factory_pos": d["factory_pos"],
        "pos_u": d["pos_u"],
        "asp": asp,
        "build": build,
        "share": share,
    }


def _calculate_consumption_year_metrics(
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]],
    all_years: List[str],
) -> Dict[str, Dict[str, Any]]:
    year_metrics = {}
    for yr in all_years:
        d = data_by_year[yr]
        prev_yr_str = str(int(yr) - 1) if yr.isdigit() else None
        prev_yr_dec_pos = None
        prev_yr_dec_asp = None
        if prev_yr_str and prev_yr_str in data_by_year:
            prev_yr_dec_pos = data_by_year[prev_yr_str]["pos_val"][11]
            p_v = data_by_year[prev_yr_str]["pos_val"][11]
            p_u = data_by_year[prev_yr_str]["pos_u"][11]
            if p_v is not None and p_u is not None and p_u != 0:
                prev_yr_dec_asp = p_v / p_u
        year_metrics[yr] = _calculate_single_year_consumption_metrics(d, prev_yr_dec_pos, prev_yr_dec_asp)
    return year_metrics


def _detect_consumption_build_outliers(
    year_metrics: Dict[str, Dict[str, Any]],
    latest_year: str,
    hist_years: List[str],
) -> Dict[int, bool]:
    red_outliers = {}
    for m_idx in range(12):
        cur_b = year_metrics[latest_year]["build"][m_idx]
        if cur_b is not None:
            hist_b = [year_metrics[h_yr]["build"][m_idx] for h_yr in hist_years if year_metrics.get(h_yr, {}).get("build", [None] * 12)[m_idx] is not None]
            if len(hist_b) >= 2:
                h_mean = statistics.mean(hist_b)
                h_std = statistics.stdev(hist_b)
                if abs(cur_b - h_mean) > (2.0 * h_std):
                    red_outliers[m_idx] = True
    return red_outliers


def build_consumption_matrix(rows_sf: List[Dict[str, Any]], model_name: str | list[str]) -> Dict[str, Any]:
    """
    Computes complete Consumption Validation matrix from Snowflake weekly data and live DB building blocks.
    Dynamically rolls years based on active fiscal year.
    """
    if not rows_sf:
        return {}

    data_by_year = _aggregate_consumption_sf_rows(rows_sf)

    curr_kv_year = get_current_kv_year()
    # Dynamic rolling 5-year window: e.g. 2022..2026 for 2026, or 2023..2027 for 2027
    min_year = curr_kv_year - 4
    all_years = sorted([yr for yr in data_by_year if yr.isdigit() and int(yr) >= min_year])
    if not all_years:
        all_years = sorted([yr for yr in data_by_year if yr.isdigit()])
    if not all_years:
        return {}

    latest_year = all_years[-1]
    prev_year = str(int(latest_year) - 1)
    latest_year_int = int(latest_year)
    hist_years = [str(latest_year_int - 5 + i) for i in range(5) if str(latest_year_int - 5 + i) in data_by_year]

    # Calculate Factory POS using Price Index multipliers
    for yr_k in all_years:
        for m_i in range(12):
            pv_val = data_by_year[yr_k]["pos_val"][m_i]
            if pv_val is not None:
                f_pos, _ = get_factory_pos_val(m_i + 1, model_name, pv_val)
                data_by_year[yr_k]["factory_pos"][m_i] = f_pos

    # Completeness check & nullify incomplete months
    comp_status = get_kv_month_completeness_status(rows_sf, target_year=latest_year)
    for yr_k in all_years:
        yr_comp = get_kv_month_completeness_status(rows_sf, target_year=yr_k)
        yr_m_details = yr_comp.get("month_details", {})
        for m_n in range(1, 13):
            if not yr_m_details.get(m_n, {}).get("is_complete", False):
                m_i = m_n - 1
                data_by_year[yr_k]["pos_val"][m_i] = None
                data_by_year[yr_k]["pos_u"][m_i] = None
                data_by_year[yr_k]["factory_pos"][m_i] = None

    _project_consumption_future_months(
        data_by_year, latest_year, prev_year, latest_year_int, comp_status["latest_complete_m_nbr"], model_name, comp_status
    )
    year_metrics = _calculate_consumption_year_metrics(data_by_year, all_years)
    red_outliers = _detect_consumption_build_outliers(year_metrics, latest_year, hist_years)

    return {
        "all_years": all_years,
        "latest_year": latest_year,
        "prev_year": prev_year,
        "year_metrics": year_metrics,
        "comp_status": comp_status,
        "red_outliers": red_outliers,
        "data_by_year": data_by_year,
    }


def _is_valid_numeric_val(val: Any) -> bool:
    return val is not None and not (isinstance(val, float) and (math.isnan(val) or math.isinf(val)))


def _accumulate_shipment_row(
    r: Dict[str, Any], data_by_year: Dict[str, Dict[str, List[Optional[float]]]]
) -> None:
    yr = str(r.get("YEAR") or 0)
    m_nbr = int(r.get("MONTHNUMBER") or 0)
    curr_kv_year = get_current_kv_year()
    # Dynamic rolling floor
    min_year = curr_kv_year - 4
    if not yr.isdigit() or int(yr) < min_year or m_nbr < 1 or m_nbr > 12:
        return
    m_idx = m_nbr - 1
    if yr not in data_by_year:
        data_by_year[yr] = {"gts_val": [None] * 12, "gts_qty": [None] * 12}
    amt_raw = r.get("GTS_AMT")
    qty_raw = r.get("GTS_QTY")
    if _is_valid_numeric_val(amt_raw):
        data_by_year[yr]["gts_val"][m_idx] = (data_by_year[yr]["gts_val"][m_idx] or 0.0) + float(amt_raw)
    if _is_valid_numeric_val(qty_raw):
        q_num = float(qty_raw)
        data_by_year[yr]["gts_qty"][m_idx] = (data_by_year[yr]["gts_qty"][m_idx] or 0.0) + (abs(q_num) if q_num != 0 else 0.0)


def _aggregate_shipment_sf_rows(rows_ship: List[Dict[str, Any]]) -> Dict[str, Dict[str, List[Optional[float]]]]:
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]] = {}
    for r in rows_ship:
        _accumulate_shipment_row(r, data_by_year)
    return data_by_year


def _compute_prev_year_fy_price_factor(
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]],
    prev_year: str,
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
) -> float:
    if prev_year not in data_by_year:
        return 1.0
    prev_usd_arr = [v for v in data_by_year[prev_year]["gts_val"] if v is not None and not (isinstance(v, float) and math.isnan(v))]
    prev_qty_arr = [v for v in data_by_year[prev_year]["gts_qty"] if v is not None and not (isinstance(v, float) and math.isnan(v))]
    prev_pos_val_arr = [pos_val_map.get((str(prev_year), i)) for i in range(12) if pos_val_map.get((str(prev_year), i)) is not None]
    prev_pos_u_arr = [pos_u_map.get((str(prev_year), i)) for i in range(12) if pos_u_map.get((str(prev_year), i)) is not None]

    tot_prev_usd = sum(prev_usd_arr) if prev_usd_arr else 0.0
    tot_prev_qty = sum(prev_qty_arr) if prev_qty_arr else 0.0
    tot_prev_pos_val = sum(prev_pos_val_arr) if prev_pos_val_arr else 0.0
    tot_prev_pos_u = sum(prev_pos_u_arr) if prev_pos_u_arr else 0.0

    asp_prev_fy = (tot_prev_pos_val / tot_prev_pos_u) if (tot_prev_pos_val > 0 and tot_prev_pos_u > 0) else None
    b3_prev_fy = (tot_prev_usd / tot_prev_qty) if (tot_prev_usd > 0 and tot_prev_qty > 0) else None
    if asp_prev_fy and b3_prev_fy and b3_prev_fy != 0:
        return asp_prev_fy / b3_prev_fy
    return 1.0


def _resolve_month_factory_pos(
    latest_year: str,
    prev_year: str,
    m_i: int,
    m_nbr: int,
    model_name: str | list[str],
    factory_pos_map: Dict[Tuple[str, int], float],
    pos_val_map: Dict[Tuple[str, int], float],
) -> float:
    factory_pos = factory_pos_map.get((str(latest_year), m_i))
    if factory_pos is not None:
        return factory_pos

    prev_pos_val = pos_val_map.get((str(prev_year), m_i))
    if prev_pos_val is not None:
        f_pos_calc, _ = get_factory_pos_val(m_nbr, model_name, prev_pos_val)
        factory_pos = f_pos_calc if f_pos_calc is not None else prev_pos_val
    else:
        factory_pos = 0.0

    factory_pos_map[(str(latest_year), m_i)] = factory_pos
    return factory_pos


def _resolve_month_asp(
    latest_year: str,
    prev_year: str,
    m_i: int,
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
) -> Optional[float]:
    p_val_m = pos_val_map.get((str(latest_year), m_i))
    p_u_m = pos_u_map.get((str(latest_year), m_i))
    if p_val_m and p_u_m and p_u_m != 0:
        return p_val_m / p_u_m

    if prev_year:
        p_val_prev = pos_val_map.get((str(prev_year), m_i))
        p_u_prev = pos_u_map.get((str(prev_year), m_i))
        if p_val_prev and p_u_prev and p_u_prev != 0:
            return p_val_prev / p_u_prev

    return None


def _project_shipment_future_months(
    data_by_year: Dict[str, Dict[str, List[Optional[float]]]],
    latest_year: str,
    prev_year: str,
    latest_comp_m_nbr: int,
    factory_pos_map: Dict[Tuple[str, int], float],
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
    fy_price_factor_prev: float,
    model_name: str | list[str],
    bb_rows: Optional[List[Dict[str, Any]]],
):
    if latest_comp_m_nbr >= 12:
        return
    if bb_rows is None:
        try:
            bb_rows = load_live_building_blocks(domain="SHIPMENT")
        except (psycopg2.Error, OSError, ValueError, KeyError):
            bb_rows = []

    for m_i in range(latest_comp_m_nbr, 12):
        m_nbr = m_i + 1
        factory_pos = _resolve_month_factory_pos(
            latest_year, prev_year, m_i, m_nbr, model_name, factory_pos_map, pos_val_map
        )
        bb_total_k = get_shipment_building_block_total_for_period(model_name, int(latest_year), m_nbr, bb_rows)
        calc_grs_usd = factory_pos + ((bb_total_k * 1000.0) if bb_total_k else 0.0)
        data_by_year[latest_year]["gts_val"][m_i] = calc_grs_usd

        asp_m = _resolve_month_asp(latest_year, prev_year, m_i, pos_val_map, pos_u_map)
        b3_proj = (asp_m / fy_price_factor_prev) if (asp_m and fy_price_factor_prev != 0) else None
        if calc_grs_usd is not None and b3_proj and b3_proj != 0:
            data_by_year[latest_year]["gts_qty"][m_i] = calc_grs_usd / b3_proj


def _calc_single_yr_b3(
    gts_val: List[Optional[float]],
    gts_qty: List[Optional[float]],
    yr: str,
    prev_year: str,
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
    fy_price_factor_prev: float,
) -> List[Optional[float]]:
    b3_arr = [None] * 12
    for m_i in range(12):
        gs = gts_val[m_i]
        gu = gts_qty[m_i]
        if gs is not None and gu is not None and gu != 0:
            b3_arr[m_i] = gs / gu
        else:
            asp_m = _resolve_month_asp(yr, prev_year, m_i, pos_val_map, pos_u_map)
            b3_arr[m_i] = (asp_m / fy_price_factor_prev) if (asp_m is not None and fy_price_factor_prev != 0) else None
    return b3_arr


def _calc_single_yr_build_bleed(
    gts_val: List[Optional[float]],
    yr: str,
    factory_pos_map: Dict[Tuple[str, int], float],
) -> List[Optional[float]]:
    build_bleed_arr = [None] * 12
    for m_i in range(12):
        gs = gts_val[m_i]
        fp = factory_pos_map.get((str(yr), m_i))
        if gs is not None and not (isinstance(gs, float) and math.isnan(gs)):
            build_bleed_arr[m_i] = (gs - fp) if (fp is not None and not (isinstance(fp, float) and math.isnan(fp))) else gs
        elif fp is not None and not (isinstance(fp, float) and math.isnan(fp)):
            build_bleed_arr[m_i] = 0.0 - fp
        else:
            build_bleed_arr[m_i] = 0.0
    return build_bleed_arr


def _calc_single_yr_unit_ratio(
    gts_val: List[Optional[float]],
    gts_qty: List[Optional[float]],
    yr: str,
    prev_year: str,
    pos_u_map: Dict[Tuple[str, int], float],
) -> List[Optional[float]]:
    unit_ratio_arr = [None] * 12
    for m_i in range(12):
        gs = gts_val[m_i]
        if gs is None or gs == 0 or (isinstance(gs, float) and math.isnan(gs)):
            unit_ratio_arr[m_i] = None
        else:
            gu = gts_qty[m_i]
            pu = pos_u_map.get((str(yr), m_i))
            if (pu is None or pu == 0) and prev_year:
                pu = pos_u_map.get((str(prev_year), m_i))
            unit_ratio_arr[m_i] = ((gu / pu) * 100.0) if (gu is not None and pu is not None and pu != 0) else None
    return unit_ratio_arr


def _calc_single_yr_price_factor(
    gts_val: List[Optional[float]],
    gts_qty: List[Optional[float]],
    yr: str,
    prev_year: str,
    b3_arr: List[Optional[float]],
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
    fy_price_factor_prev: float,
) -> List[Optional[float]]:
    price_factor_arr = [None] * 12
    for m_i in range(12):
        gs = gts_val[m_i]
        if gs is None or gs == 0 or (isinstance(gs, float) and math.isnan(gs)):
            price_factor_arr[m_i] = None
        else:
            asp_m = _resolve_month_asp(yr, prev_year, m_i, pos_val_map, pos_u_map)
            b3_m = b3_arr[m_i]
            if b3_m is None and gts_qty[m_i] is not None and gts_qty[m_i] != 0:
                b3_m = gs / gts_qty[m_i]
            price_factor_arr[m_i] = (asp_m / b3_m) if (asp_m is not None and b3_m is not None and b3_m != 0) else fy_price_factor_prev
    return price_factor_arr


def _calculate_shipment_single_year_metrics(
    yr: str,
    d: Dict[str, List[Optional[float]]],
    prev_year: str,
    factory_pos_map: Dict[Tuple[str, int], float],
    pos_val_map: Dict[Tuple[str, int], float],
    pos_u_map: Dict[Tuple[str, int], float],
    fy_price_factor_prev: float,
) -> Dict[str, Any]:
    gts_val = d["gts_val"]
    gts_qty = d["gts_qty"]

    b3_arr = _calc_single_yr_b3(gts_val, gts_qty, yr, prev_year, pos_val_map, pos_u_map, fy_price_factor_prev)
    build_bleed_arr = _calc_single_yr_build_bleed(gts_val, yr, factory_pos_map)
    unit_ratio_arr = _calc_single_yr_unit_ratio(gts_val, gts_qty, yr, prev_year, pos_u_map)
    price_factor_arr = _calc_single_yr_price_factor(gts_val, gts_qty, yr, prev_year, b3_arr, pos_val_map, pos_u_map, fy_price_factor_prev)

    return {
        "gts_val": gts_val,
        "gts_qty": gts_qty,
        "b3": b3_arr,
        "build_bleed": build_bleed_arr,
        "unit_ratio": unit_ratio_arr,
        "price_factor": price_factor_arr,
        "gross_ship": gts_val,
        "gross_case": gts_qty,
        "net_ship": b3_arr,
        "share": unit_ratio_arr,
    }


def _evaluate_month_unit_ratio_outlier(
    m_i: int,
    year_metrics: Dict[str, Dict[str, Any]],
    latest_year: str,
    ref_years: List[str],
) -> Optional[str]:
    v_curr = year_metrics[latest_year]["unit_ratio"][m_i]
    if v_curr is None or (isinstance(v_curr, float) and math.isnan(v_curr)):
        return None

    ref_4_vals = [
        year_metrics[y]["unit_ratio"][m_i]
        for y in ref_years
        if y in year_metrics
        and year_metrics[y]["unit_ratio"][m_i] is not None
        and not (isinstance(year_metrics[y]["unit_ratio"][m_i], float) and math.isnan(year_metrics[y]["unit_ratio"][m_i]))
    ]
    if len(ref_4_vals) < 4:
        return None

    sum_ref = sum(ref_4_vals)
    baseline = (sum_ref - min(ref_4_vals) - max(ref_4_vals)) / 2.0
    if baseline <= 0:
        return None

    rel_diff_pct = abs((v_curr - baseline) / baseline) * 100.0
    if rel_diff_pct > 10.0:
        return "red"
    if rel_diff_pct > 5.0:
        return "yellow"
    return None


def _detect_unit_ratio_outliers(
    year_metrics: Dict[str, Dict[str, Any]],
    latest_year: str,
    ref_years: List[str],
) -> Dict[int, str]:
    unit_ratio_outliers = {}
    if latest_year not in year_metrics:
        return unit_ratio_outliers

    for m_i in range(12):
        color = _evaluate_month_unit_ratio_outlier(m_i, year_metrics, latest_year, ref_years)
        if color:
            unit_ratio_outliers[m_i] = color

    return unit_ratio_outliers


def _detect_shipment_cutoff_month(latest_year: str, latest_gts_vals: List[Optional[float]]) -> int:
    """Detects latest completed shipment month dynamically. Returns 0 if no actuals yet."""
    for m_i in range(11, -1, -1):
        v = latest_gts_vals[m_i]
        if v is not None and not (isinstance(v, float) and math.isnan(v)) and v > 0:
            return m_i + 1
    return 0


def _blend_factory_pos_override(
    factory_pos_map: Dict[Tuple[str, int], float],
    factory_pos_by_year: Optional[Dict[str, List[Optional[float]]]],
) -> None:
    if not factory_pos_by_year:
        return
    for yr_k, fp_list in factory_pos_by_year.items():
        for m_i, val in enumerate(fp_list):
            if val is not None:
                factory_pos_map[(str(yr_k), m_i)] = val


def build_shipment_matrix(
    rows_ship: List[Dict[str, Any]],
    model_name: str | list[str],
    factory_pos_by_year: Optional[Dict[str, List[Optional[float]]]] = None,
    rows_cons: Optional[List[Dict[str, Any]]] = None,
    bb_rows: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """
    Computes complete Shipment Validation matrix matching validation/ship.py specifications.
    Dynamically rolls years based on active fiscal year.
    """
    if not rows_ship:
        return {}

    data_by_year = _aggregate_shipment_sf_rows(rows_ship)

    curr_kv_year = get_current_kv_year()
    min_year = curr_kv_year - 4
    all_years = sorted([yr for yr in data_by_year if yr.isdigit() and int(yr) >= min_year])
    if not all_years:
        all_years = sorted([yr for yr in data_by_year if yr.isdigit()])
    if not all_years:
        return {}

    latest_year = all_years[-1]
    prev_year = str(int(latest_year) - 1)
    ref_years = [y for y in all_years if y < latest_year][-4:]

    factory_pos_map, pos_val_map, pos_u_map = extract_consumption_maps(rows_cons, model_name) if rows_cons else ({}, {}, {})
    _blend_factory_pos_override(factory_pos_map, factory_pos_by_year)

    fy_price_factor_prev = _compute_prev_year_fy_price_factor(data_by_year, prev_year, pos_val_map, pos_u_map)
    latest_comp_m_nbr = _detect_shipment_cutoff_month(latest_year, data_by_year[latest_year]["gts_val"])

    _project_shipment_future_months(
        data_by_year, latest_year, prev_year, latest_comp_m_nbr, factory_pos_map, pos_val_map, pos_u_map, fy_price_factor_prev, model_name, bb_rows
    )

    year_metrics = {
        yr: _calculate_shipment_single_year_metrics(yr, data_by_year[yr], prev_year, factory_pos_map, pos_val_map, pos_u_map, fy_price_factor_prev)
        for yr in all_years
    }
    unit_ratio_outliers = _detect_unit_ratio_outliers(year_metrics, latest_year, ref_years)

    comp_status = {
        "latest_complete_m_nbr": latest_comp_m_nbr,
        "ytd_slice": [0, latest_comp_m_nbr],
        "ytg_slice": [latest_comp_m_nbr, 12],
    }

    return {
        "all_years": all_years,
        "latest_year": latest_year,
        "prev_year": prev_year,
        "year_metrics": year_metrics,
        "comp_status": comp_status,
        "unit_ratio_outliers": unit_ratio_outliers,
        "red_outliers": {},
        "data_by_year": data_by_year,
        "pos_val_map": pos_val_map,
        "pos_u_map": pos_u_map,
        "factory_pos_map": factory_pos_map,
        "fy_price_factor_prev": fy_price_factor_prev,
    }
